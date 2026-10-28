"""Spreadsheet import: CSV / TSV / XLSX -> contact rows.

Two steps so the user can check the column mapping before anything is saved:
`parse` returns the headers, a guessed mapping and a preview; the router then
calls `to_contacts` with the (possibly corrected) mapping.

Every column is kept: the ones with a field of their own (below) are mapped,
and any other column goes into the contact's `extra` notes. An Excel file is
read tab by tab (e.g. one tab per city, "B_HYD", "B_MUM"); tabs whose header
row isn't recognised (a "Read Me") are left out. Cells that only say a value
is missing ("Not publicly listed", "N/A") count as empty.
"""
import csv
import io
import json
import re

from . import marketplace

# "type" comes before "category" so a file with only a Category column keeps using it as the type.
FIELDS = ["name", "email", "type", "category", "specialty", "clinic", "product", "city", "notes",
          "description", "notes_source", "partner_id", "tier", "phone", "website", "socials", "pricing",
          "spend", "spend_source", "source_url", "needs_check"]

SYNONYMS = {
    "name": ["name", "full name", "doctor", "doctor name", "contact", "contact name", "person"],
    "email": ["email", "e-mail", "email address", "mail", "email id"],
    "type": ["type", "contact type", "category", "segment"],
    "specialty": ["specialty", "speciality", "specialization", "specialisation", "department", "role", "designation", "title", "position"],
    "clinic": ["clinic", "hospital", "clinic / hospital", "hospital / clinic", "clinic/hospital", "hospital/clinic", "clinic name", "hospital name", "company name", "organisation", "organization", "company", "practice", "institution"],
    "product": ["product", "product name", "device", "brand", "offering", "hero product", "hero products", "key product"],
    "city": ["city", "location", "town", "district"],
    "notes": ["notes", "note", "comments", "remarks", "context", "personalisation", "personalization"],
    # The brand sheet calls it "Category Type" or "B-Category" (and, on one tab, just "FD").
    "category": ["category", "sub category", "subcategory", "category type", "b-category", "b category",
                 "p-category", "p category", "brand category", "fd"],
    "description": ["description", "about", "what they do"],
    # The founder's sheet: one free-text cell holding site, description, source and a CHECK note.
    "notes_source": ["notes / source", "notes/source", "notes & source", "notes and source"],
    "partner_id": ["id", "partner id", "partner id."],
    "tier": ["tier"],
    "phone": ["phone", "mobile", "phone number", "contact number", "whatsapp"],
    "website": ["website", "site", "web", "url"],
    "socials": ["instagram / socials", "instagram/socials", "socials", "social", "social media", "social links",
                "social profiles", "instagram", "linkedin", "facebook"],
    # Kept for the team only: the drafter never sees pricing.
    "pricing": ["pricing", "price", "prices", "price range"],
    # Angle A of the brand guide needs both, filled in by hand: never estimated.
    "spend": ["spend last quarter (₹ lakhs)", "spend last quarter (rs lakhs)", "spend last quarter", "marketing spend", "spend"],
    "spend_source": ["spend source", "source of spend"],
    "source_url": ["source", "source url"],
    "needs_check": ["check", "needs check", "to check"],
}
FIRST_LAST = (["first name", "firstname", "first"], ["last name", "lastname", "surname", "last"])
# Row numbers and the like: not data about the partner, so not kept in `extra`.
IGNORED = {"s.no.", "s.no", "s no", "s. no.", "sr no", "sr. no.", "sr.no.", "serial", "serial no", "#", "no", "no.", "sl no", "sl. no."}
# Placeholders a research sheet uses instead of leaving a cell blank.
MISSING_RE = re.compile(r"^(not (publicly )?(listed|available|found)|n/?a|na|nil|none|null|-+|—|–|\?|unknown|tbd)\.?$", re.I)

TYPE_VALUES = {
    "doctor": ["doctor", "dr", "specialist", "physician", "clinic", "clinician", "surgeon"],
    "hospital": ["hospital", "institution", "institutional", "hospital pilot"],
    "company": ["company", "network", "insurer", "diagnostics", "chain", "corporate"],
    "partner": ["partner", "product", "product partner", "device", "vendor"],
    "brand": ["brand", "b", "health brand", "marketplace brand"],
    "professional": ["professional", "p", "health professional", "nutritionist", "dietitian", "physio", "trainer", "therapist", "counsellor", "psychologist"],
}
CITY_CODES = {"HYD": "Hyderabad", "DEL": "Delhi", "MUM": "Mumbai", "BLR": "Bengaluru"}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
# Brand rows that aren't brands you can buy through the chat (the brand guide's skip rules).
FOOD_DELIVERY_RE = re.compile(r"(^|[/.])(zomato|swiggy)\.com\b", re.I)


class ImportError_(Exception):
    pass


def _norm(h: str) -> str:
    return re.sub(r"[\s_]+", " ", (h or "").strip().lower())


def blank_if_missing(value: str) -> str:
    v = (value or "").strip()
    return "" if MISSING_RE.match(v) else v


def canonical(header: str) -> str:
    """The field a header belongs to, or the header itself, so tabs that name a column differently line up."""
    h = _norm(header)
    for field, words in SYNONYMS.items():
        if h in words:
            return field
    return h


def tab_city(tab: str) -> str:
    """'B_HYD' -> 'Hyderabad', 'Mumbai' -> 'Mumbai', 'Brands' -> ''."""
    t = (tab or "").strip()
    m = re.fullmatch(r"(?:[BP][_\- ])?([A-Za-z]{3})", t)
    if m and m.group(1).upper() in CITY_CODES:
        return CITY_CODES[m.group(1).upper()]
    return next((c for c in CITY_CODES.values() if c.lower() == t.lower()), "")


def read_tables(filename: str, data: bytes) -> dict:
    """{'headers', 'rows', 'row_tabs', 'tabs': [{'name', 'rows'}], 'skipped_tabs'}. A CSV is one tab named ''.
    Tabs are lined up by what each column means, so "Category Type" on one tab and "FD" on another become one column."""
    name = (filename or "").lower()
    if not name.endswith((".xlsx", ".xlsm")):
        headers, rows = read_table(filename, data)
        return {"headers": headers, "rows": rows, "row_tabs": [""] * len(rows),
                "tabs": [{"name": "", "rows": len(rows)}], "skipped_tabs": []}
    wb = _open_workbook(data)
    found, skipped = [], []
    for ws in wb.worksheets:
        table = _clean([["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)])
        start = _find_header(table)
        if start is None:
            skipped.append(ws.title)
            continue
        found.append((ws.title, table[start], table[start + 1:]))
    wb.close()
    if not found:
        headers, rows = read_table(filename, data)  # no recognisable headers anywhere: first tab, first row
        return {"headers": headers, "rows": rows, "row_tabs": [""] * len(rows),
                "tabs": [{"name": "", "rows": len(rows)}], "skipped_tabs": []}
    keys, headers = [], []
    for _, hdr, _ in found:
        for h in hdr:
            k = canonical(h)
            if h and k not in keys:
                keys.append(k)
                headers.append(h)
    rows, row_tabs, tabs = [], [], []
    for title, hdr, body in found:
        col = {canonical(h): i for i, h in reversed(list(enumerate(hdr))) if h}
        for r in body:
            rows.append([(r[col[k]] if k in col and col[k] < len(r) else "") for k in keys])
            row_tabs.append(title)
        tabs.append({"name": title, "rows": len(body)})
    return {"headers": headers, "rows": rows, "row_tabs": row_tabs, "tabs": tabs, "skipped_tabs": skipped}


def _open_workbook(data: bytes):
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise ImportError_("XLSX support needs openpyxl installed.") from e
    try:
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ImportError_(f"Couldn't read that workbook: {e}") from e


def _clean(table: list[list[str]]) -> list[list[str]]:
    return [r for r in table if any(c for c in r)]


def read_table(filename: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = _open_workbook(data)
        ws = wb.worksheets[0]
        table = [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
    elif name.endswith(".xls"):
        raise ImportError_("Old .xls files aren't supported — save it as .xlsx or .csv first.")
    else:
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel_tab if "\t" in text.splitlines()[0] else csv.excel
        table = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    table = [r for r in table if any(c for c in r)]
    if not table:
        raise ImportError_("That file is empty.")
    start = header_row(table)
    headers, body = table[start], table[start + 1:]
    width = len(headers)
    body = [(r + [""] * width)[:width] for r in body]
    return headers, body


def header_row(table: list[list[str]]) -> int:
    """Index of the header row. Usually 0, but the founder's tabs start with a title and a note."""
    found = _find_header(table)
    return 0 if found is None else found


def _find_header(table: list[list[str]]) -> int | None:
    known = {w for words in SYNONYMS.values() for w in words}
    for i, row in enumerate(table[:10]):
        if sum(1 for c in row if _norm(c) in known) >= 3:
            return i
    return None


def guess_mapping(headers: list[str], rows: list[list[str]]) -> dict:
    """field -> column index (or None). First/last-name columns map to name_first/name_last."""
    normed = [_norm(h) for h in headers]
    mapping: dict = {f: None for f in FIELDS}

    def looks_like_email(i):
        sample = [r[i] for r in rows[:50] if r[i]]
        return sample and sum(1 for v in sample if EMAIL_RE.match(v)) / len(sample) >= 0.6

    # Email first: by header, else the column whose values are mostly addresses.
    mapping["email"] = next((i for i, h in enumerate(normed) if h in SYNONYMS["email"]), None)
    if mapping["email"] is None:
        mapping["email"] = next((i for i in range(len(headers)) if looks_like_email(i)), None)
    used = {mapping["email"]} - {None}
    for field in FIELDS:
        if field == "email":
            continue
        for i, h in enumerate(normed):
            if i not in used and h in SYNONYMS[field] and not looks_like_email(i):
                mapping[field] = i
                used.add(i)
                break
    if mapping["name"] is None:
        first = next((i for i, h in enumerate(normed) if h in FIRST_LAST[0]), None)
        last = next((i for i, h in enumerate(normed) if h in FIRST_LAST[1]), None)
        if first is not None:
            mapping["name_first"], mapping["name_last"] = first, last
    return mapping


def normalise_type(value: str, default: str) -> str:
    v = _norm(value)
    if not v:
        return default
    for t, words in TYPE_VALUES.items():
        if v == t or v in words:
            return t
    for t, words in TYPE_VALUES.items():
        if any(len(w) > 1 and w in v for w in words):
            return t
    return default


def first_email(cell: str) -> str:
    """'contact@jiva.com; clinics@jiva.com' -> 'contact@jiva.com' (the brand guide: use the first one)."""
    for part in re.split(r"[\s,;/]+", cell or ""):
        part = part.strip().strip("<>()").lower()
        if part:
            return part
    return ""


def brand_key(name: str) -> str:
    """Same brand, written differently: 'Fast&Up (Aeronutrix)' and 'FAST&UP Pvt Ltd' -> 'fastup'."""
    n = marketplace.shop_name(name)
    n = re.sub(r"\b(pvt\.?|private|ltd\.?|limited|llp|inc\.?|india)\b", " ", n, flags=re.I)
    return re.sub(r"[^a-z0-9]", "", n.lower())


def to_contacts(headers: list[str], rows: list[list[str]], mapping: dict, default_type: str = "doctor",
                audience: str = "", row_tabs: list[str] | None = None) -> tuple[list[dict], list[dict]]:
    """Returns (contacts, problems). Problems are rows that couldn't become a contact.
    `audience` ('brand' or 'professional', chosen on upload) is the type of every row, whatever a Type column says."""
    def cell(row, field):
        i = mapping.get(field)
        return blank_if_missing(row[int(i)]) if i is not None and i != "" and int(i) < len(row) else ""

    used = {int(i) for f, i in mapping.items() if i is not None and i != ""}
    extra_cols = [(i, h) for i, h in enumerate(headers) if i not in used and h.strip() and _norm(h) not in IGNORED]
    contacts, problems = [], []
    seen_brands: dict[str, str] = {}
    tabs = row_tabs or [""] * len(rows)
    numbers: dict[str, int] = {}
    for row, tab in zip(rows, tabs):
        n = numbers[tab] = numbers.get(tab, 1) + 1  # row 1 is the header
        where = {"row": n, **({"tab": tab} if tab else {})}
        name = cell(row, "name")
        if not name and mapping.get("name_first") is not None:
            name = " ".join(x for x in (cell(row, "name_first"), cell(row, "name_last")) if x)
        email = first_email(cell(row, "email"))
        if not email:
            problems.append({**where, "name": name, "reason": "no email"})
            continue
        if not EMAIL_RE.match(email):
            problems.append({**where, "name": name, "reason": f"invalid email {email!r}"})
            continue
        notes, remarks_id = marketplace.split_partner_id(cell(row, "notes"))
        partner_id = cell(row, "partner_id") or remarks_id
        from_id = marketplace.parse_partner_id(partner_id)
        if audience in marketplace.TYPES:
            ctype = audience
        else:
            ctype = normalise_type(cell(row, "type"), "") or from_id.get("type") or default_type
        parsed = marketplace.parse_notes_source(cell(row, "notes_source"))
        city = cell(row, "city") or tab_city(tab)
        market = ctype in marketplace.TYPES
        website = cell(row, "website") or parsed["website"]
        if ctype == "brand" and FOOD_DELIVERY_RE.search(website):
            problems.append({**where, "name": name, "reason": "a Zomato or Swiggy page (a restaurant, not a brand sold in the chat)"})
            continue
        if ctype == "brand" and (key := brand_key(name)):
            if key in seen_brands:
                problems.append({**where, "name": name, "reason": f"the same brand as {seen_brands[key]}"})
                continue
            seen_brands[key] = f"{tab} row {n}" if tab else f"row {n}"
        extra = {h.strip(): v for i, h in extra_cols if i < len(row) and (v := blank_if_missing(row[i]))}
        contacts.append({
            "name": name or email.split("@")[0],
            "email": email,
            "type": ctype,
            # Marketplace rows: the category may sit in its own column, in Speciality, or only in the partner ID.
            "category": (marketplace.guess_category(ctype, cell(row, "category"), cell(row, "specialty"))
                         or from_id.get("category", "")) if market else "",
            "specialty": cell(row, "specialty"),
            "clinic": cell(row, "clinic"),
            "product": cell(row, "product"),
            "city": CITY_CODES.get(city.upper(), city),
            "notes": notes,
            "partner_id": partner_id,
            "tier": cell(row, "tier").upper().replace("TBD", "") or from_id.get("tier", ""),
            "phone": cell(row, "phone"),
            **{f: cell(row, f) or parsed[f] for f in ("description", "source_url", "needs_check")},
            "website": website,
            **{f: cell(row, f) for f in ("socials", "pricing", "spend", "spend_source")},
            "extra": json.dumps(extra, ensure_ascii=False) if extra else "",
        })
        if market and not contacts[-1]["description"]:
            # The simple CSV layout puts a marketplace partner's one-line description in Product.
            contacts[-1]["description"] = contacts[-1]["product"]
    return contacts, problems
