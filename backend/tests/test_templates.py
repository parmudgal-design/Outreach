"""Brand / professional uploads, per-audience templates, the banner in the email body, and research
from the website and socials. All contacts are fake."""
import io
import json
from types import SimpleNamespace

import openpyxl
import pytest

from backend.app import db, worker
from backend.app.services import drafting, importer, mailer, marketplace, outreach, prompts, qa, research, spend, templates
from backend.app.settings_store import DEFAULTS

SENDER = {"name": "Ananya Gupta", "role": "Founder, Dr.Tracker", "tone": "warm", "address": "HSR Layout, Bengaluru 560102"}
MARKET = {**DEFAULTS["marketplace"], "meeting_link": "https://cal.example/drtracker"}
SPEND = {"spend": "18-25", "spend_source": "ad library + influencer post estimates, Jul–Sep 2026"}
BRAND_HEADERS = ["S.No.", "Name", "B-Category", "Hero Product", "Email", "Website", "Instagram / Socials", "Pricing",
                 "Phone Number", "Founded"]


def brand_workbook() -> bytes:
    """The brand sheet's shape: one tab per city (B_HYD calls the category column "FD"), a test tab and a Read Me."""
    wb = openpyxl.Workbook()
    hyd = wb.active
    hyd.title = "B_HYD"
    hyd.append(["S.No.", "Name", "FD", "Hero Product", "Email", "Website", "Instagram / Socials", "Pricing", "Phone Number", "Founded"])
    hyd.append([1, "Heka Bites Pvt Ltd", "FD", "Premium Raw Makhana 250g / Peri Peri Makhana", "hello@heka.test; sales@heka.test",
                "https://heka.test", "@hekabites", "₹199–₹499", "98765 43210", "2019"])
    hyd.append([2, "Biryani Box", "FD", "Chicken Biryani", "orders@biryanibox.test", "https://www.zomato.com/hyderabad/biryani-box",
                "Not publicly listed", "Not publicly listed", "", ""])
    hyd.append([3, "No Mail Foods", "FD", "Millet Flakes", "Not publicly listed", "https://nomail.test", "", "", "", ""])
    mum = wb.create_sheet("B_MUM")
    mum.append(BRAND_HEADERS)
    mum.append([1, "GoodShift (Mumbai Hub)", "PR", "Whey Protein Chocolate Supreme; Protein Bars", "team@goodshift.test",
                "goodshift.test", "https://www.instagram.com/goodshift/ https://in.linkedin.com/company/goodshift", "₹1,999", "", ""])
    mum.append([2, "HEKA BITES PRIVATE LIMITED", "FD", "Makhana", "mumbai@heka.test", "", "", "", "", ""])  # same brand
    mum.append([3, "OZiva", "SUP", "Plant-Based Nutrition", "care@oziva.test", "oziva.test", "", "", "", ""])
    test = wb.create_sheet("Brands")
    test.append(BRAND_HEADERS)
    test.append([1, "Test Brand", "DEV", "Smart Ring", "test@ringco.test", "", "", "", "", ""])
    readme = wb.create_sheet("Read Me")
    readme.append(["Item", "Details"])
    readme.append(["Total brands", "6"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def upload(c, data: bytes, audience: str, name: str = "brands.xlsx"):
    r = c.post("/api/import/preview", files={"file": (name, data)}, data={"audience": audience})
    assert r.status_code == 200, r.text
    return r.json()


def contacts(c):
    return {x["email"]: x for x in c.get("/api/contacts").json()}


# --- Upload: brands or professionals --------------------------------------------------

def test_brand_sheet_reads_every_city_tab_and_keeps_every_column(configured):
    c = configured
    p = upload(c, brand_workbook(), "brand")
    assert [t["name"] for t in p["tabs"]] == ["B_HYD", "B_MUM", "Brands"] and p["skipped_tabs"] == ["Read Me"]
    m = p["mapping"]
    assert p["headers"][m["category"]] in ("FD", "B-Category") and p["headers"][m["product"]] == "Hero Product"
    assert p["headers"][m["socials"]] == "Instagram / Socials" and p["headers"][m["pricing"]] == "Pricing"

    r = c.post("/api/import/commit", json={"token": p["token"], "mapping": m, "tabs": ["B_HYD", "B_MUM"]}).json()
    assert r["audience"] == "brand" and r["added"] == 3
    reasons = {x["name"]: x["reason"] for x in r["invalid"]}
    assert "Zomato" in reasons["Biryani Box"] and reasons["No Mail Foods"] == "no email"
    assert reasons["HEKA BITES PRIVATE LIMITED"].startswith("the same brand as B_HYD row 2")

    got = contacts(c)
    assert set(got) == {"hello@heka.test", "team@goodshift.test", "care@oziva.test"}  # test tab left out
    heka = got["hello@heka.test"]  # first of two addresses
    assert (heka["type"], heka["category"], heka["city"]) == ("brand", "FD", "Hyderabad")  # city from the tab name
    assert heka["product"].startswith("Premium Raw Makhana") and heka["socials"] == "@hekabites"
    assert heka["pricing"] == "₹199–₹499" and heka["phone"] == "98765 43210"
    assert json.loads(heka["extra"]) == {"Founded": "2019"}  # columns with no field of their own are kept; S.No. isn't
    assert got["team@goodshift.test"]["category"] == "PRO" and got["team@goodshift.test"]["city"] == "Mumbai"


def test_audience_overrides_a_type_column(configured):
    c = configured
    sheet = b"Name,Email,Type,Category\nMeenal Joshi,meenal@x.test,brand,Nutritionist\n"
    p = upload(c, sheet, "professional", "list.csv")
    c.post("/api/import/commit", json={"token": p["token"], "mapping": p["mapping"]})
    assert contacts(c)["meenal@x.test"]["type"] == "professional"
    assert c.post("/api/import/preview", files={"file": ("l.csv", sheet)}, data={"audience": "doctors"}).status_code == 400


def test_brand_already_in_the_app_is_not_added_twice(configured):
    c = configured
    c.post("/api/contacts", json={"name": "Heka Bites", "email": "old@heka.test", "type": "brand", "category": "FD"})
    p = upload(c, brand_workbook(), "brand")
    r = c.post("/api/import/commit", json={"token": p["token"], "mapping": p["mapping"], "tabs": ["B_HYD"]}).json()
    assert r["added"] == 0 and any("already in your contacts (old@heka.test)" in x["reason"] for x in r["invalid"])


def test_founder_workbook_upload_for_brands_skips_professionals(configured):
    from pathlib import Path
    data = (Path(__file__).parent / "fixtures" / "partner_list_sample.xlsx").read_bytes()
    mixed = upload(configured, data, "mixed")
    brands = upload(configured, data, "brand")
    assert brands["counts"]["add"] < mixed["counts"]["add"]
    assert brands["counts"]["skipped_by_reason"].get("not a brand")
    assert {a["type"] for a in brands["adds"]} == {"brand"}


@pytest.mark.parametrize("cell", ["Not publicly listed", "N/A", "-", "  na ", "TBD"])
def test_missing_markers_are_blank(cell):
    assert importer.blank_if_missing(cell) == ""


# --- Templates ------------------------------------------------------------------------

def test_brand_template_ships_and_professional_can_be_added(configured, tmp_path):
    c = configured
    t = c.get("/api/templates").json()
    # Brands use brand_outreach_agent.md and the three Step screenshots at the repo root, and nothing else.
    assert templates.builtin_guide_path("brand").name == "brand_outreach_agent.md"
    assert t["brand"]["guide_source"] == "built-in" and "Brand Outreach Agent" in t["brand"]["guide"]
    assert t["brand"]["banner_source"] == "built-in" and t["brand"]["banner_dims"][0] == templates.EMAIL_WIDTH
    assert [i["name"] for i in t["brand"]["images"]] == [
        "Step_1___Asks_in_chat_2x.jpg", "Step_2___Sees_it_first_2x.jpg", "Step_3___Buys_with_a_code_2x.jpg"]
    assert t["brand"]["images"][2]["alt"].startswith("Step 3 of 3:")
    assert t["brand"]["banner_size"] < 600_000  # all three, resized for email
    assert t["professional"]["guide_source"] == "" and t["professional"]["banner_source"] == ""

    for i in range(3):
        img = c.get(f"/api/templates/brand/image/{i}")
        assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg"
    assert c.get("/api/templates/brand/image/3").status_code == 404
    assert c.get("/api/templates/professional/banner.jpg").status_code == 404

    from PIL import Image
    png = io.BytesIO()
    Image.new("RGBA", (2400, 1000), (10, 120, 90, 255)).save(png, "PNG")
    t = c.post("/api/templates/professional/banner", files={"file": ("pro.png", png.getvalue())}).json()
    assert t["professional"]["banner_source"] == "custom" and t["professional"]["banner_dims"] == [1200, 500]
    t = c.put("/api/templates/professional/guide", json={"text": "# Professional guide\nWrite kindly."}).json()
    assert t["professional"]["guide_source"] == "custom"
    assert c.post("/api/templates/professional/banner", files={"file": ("x.png", b"not an image")}).status_code == 400
    t = c.delete("/api/templates/professional/banner").json()
    assert t["professional"]["banner_source"] == ""

    t = c.put("/api/templates/brand/guide", json={"text": "# My brand guide"}).json()
    assert t["brand"]["guide_source"] == "custom"
    assert c.delete("/api/templates/brand/guide").json()["brand"]["guide_source"] == "built-in"


# --- The banner in the email ------------------------------------------------------------

class Capture:
    name = "dry_run"

    def __init__(self):
        self.sent = []

    def send(self, msg):
        self.sent.append(msg)


def _send(c, cid, kind="initial", transport=None):
    m = c.post(f"/api/contacts/{cid}/draft?kind={kind}").json()
    assert not qa.has_errors(m["qa"]), m["qa"]
    c.post(f"/api/messages/{m['id']}/approve")
    transport = transport or Capture()
    outreach.send(m["id"], transport=transport)
    return transport.sent[-1]


def test_brand_first_email_carries_the_three_screenshots_inline_in_order(configured):
    c = configured
    cid = c.post("/api/contacts", json={"name": "Heka Bites", "email": "hello@heka.test", "type": "brand", "category": "FD",
                                         "product": "Premium Raw Makhana", **SPEND}).json()["id"]
    msg = _send(c, cid)
    assert msg.get_content_type() == "multipart/alternative"
    text = msg.get_body(("plain",)).get_content()
    related = msg.get_payload()[1]
    assert related.get_content_type() == "multipart/related"
    assert "[HOW_IT_WORKS_IMAGE]" not in text and "Hi Heka Bites team," in text and "unsubscribe" in text
    images = [p for p in related.iter_parts() if p.get_content_maintype() == "image"]
    assert [p.get_filename() for p in images] == [
        "Step_1___Asks_in_chat_2x.jpg", "Step_2___Sees_it_first_2x.jpg", "Step_3___Buys_with_a_code_2x.jpg"]
    page = msg.get_body(("html",)).get_content()
    assert "[SCREENSHOT_" not in page and "[SCREENSHOT_" not in text
    at = [page.index(f'src="cid:{p["Content-ID"].strip("<>")}"') for p in images]
    # Inside the body where the markers were: 1 after the hook, 2 and 3 after "But we can take…", then the call link.
    assert page.index("dance with it online") < at[0] < page.index("But we can take") < at[1] < at[2] < page.index("Set up a call")
    # In the body, not attachments: mail clients list "attachment" parts as attachments even when they're shown inline.
    assert all(p["Content-Disposition"].startswith("inline") for p in images)
    assert 'alt="Step 1 of 3:' in page and 'max-width:600px' in page
    assert all(p.get_content_type() == "image/jpeg" and len(p.get_content()) < 300_000 for p in images)

    # The follow-up goes in the same thread as plain text, without the screenshots.
    with db.connect() as conn:
        conn.execute("UPDATE contacts SET last_sent_at = 0 WHERE id = ?", (cid,))
    follow = _send(c, cid, "followup")
    assert follow.get_content_type() == "text/plain"


def test_professional_email_without_a_banner_is_plain_text_and_drops_a_marker(configured):
    c = configured
    cid = c.post("/api/contacts", json={"name": "Dt. Meenal Joshi", "email": "meenal@x.test", "type": "professional",
                                         "category": "NT"}).json()["id"]
    m = c.post(f"/api/contacts/{cid}/draft").json()
    c.put(f"/api/messages/{m['id']}", json={"body": m["body"].replace("\n\n", "\n\n[HOW_IT_WORKS_IMAGE]\n\n", 1)})
    c.post(f"/api/messages/{m['id']}/approve")
    t = Capture()
    outreach.send(m["id"], transport=t)
    assert t.sent[0].get_content_type() == "text/plain" and "HOW_IT_WORKS" not in t.sent[0].get_content()


def test_images_go_after_the_pitch_when_the_marker_was_removed():
    body = "Hi Heka team,\n\nLine one.\nLine two.\n\nWorth 15 minutes?\n\nThanks,\nAnanya"
    placed = templates.place_marker(body)
    assert placed.split("\n\n")[2] == templates.MARKER
    banner = templates.Banner(b"\xff\xd8fake", "jpeg", 1200, 700, "built-in", "alt")
    msg = mailer.build_message(from_email="a@x.test", from_name="A", to_email="b@y.test", to_name="", subject="s",
                               body=body, images=[banner])
    page = msg.get_body(("html",)).get_content()
    assert page.index("Line two") < page.index("cid:") < page.index("Worth 15 minutes")


def test_images_go_below_the_numbered_steps_when_the_marker_was_removed():
    body = "Hi Heka team,\n\nIntro.\n\nHere's how it works:\n1. Ask\n2. Browse\n3. Buy\n\nSet up a call: x\n\nAditi"
    assert templates.place_marker(body).split("\n\n")[3] == templates.MARKER


# --- Pre-send checks ----------------------------------------------------------------------

BRAND = {"name": "Heka Bites", "email": "hello@heka.test", "type": "brand", "category": "FD"}


def check(body, contact=BRAND, kind="initial", market=MARKET):
    return qa.check("Straight into your buyer's chat", body, contact, SENDER, DEFAULTS["pitches"], kind, market)


def pitch(*middle):
    return "\n\n".join(["Hi Heka Bites team,", *middle, "Worth 15 minutes with our team?", "Thanks,\nAnanya Gupta\nFounder, Dr. Tracker"])


def test_marker_is_not_a_placeholder_and_shopper_prices_are_not_pricing():
    issues = check(pitch("They compare store prices and check out in the chat.", templates.MARKER))
    assert not qa.has_errors(issues), issues


def test_marker_twice_blocks_and_missing_marker_warns():
    assert qa.has_errors(check(pitch(templates.MARKER, "More.", templates.MARKER)))
    missing = check(pitch("No image here."))
    assert not qa.has_errors(missing) and any("No image marker" in i["message"] for i in missing)


def test_spend_figure_only_as_the_sheet_has_it():
    contact = {**BRAND, "spend": "1.5–2", "spend_source": "Founder call"}
    assert not qa.has_errors(check(pitch("Last quarter, Heka Bites spent around ₹1.5–2 lakhs on marketing.", templates.MARKER), contact))
    assert qa.has_errors(check(pitch("Last quarter, Heka Bites spent around ₹3 lakhs on marketing.", templates.MARKER), contact))


def test_coupon_outside_the_trial_text_blocks():
    assert qa.has_errors(check(pitch("Use a Dr. Tracker coupon at checkout.", templates.MARKER)))


def test_unspaced_name_warns_in_guided_brand_emails():
    issues = check(pitch("Someone asks Dr.Tracker on WhatsApp for a snack.", templates.MARKER))
    assert any("without a space" in i["message"] for i in issues)


# --- Drafting to the guide -------------------------------------------------------------------

def test_keyless_brand_template_is_the_guides_worked_example():
    """brand_outreach_agent.md section 10: the same input gives the same email, screenshot markers included."""
    contact = {"name": "GreenLeaf Nutrition", "email": "rohan@greenleaf.example", "type": "brand", "category": "PRO",
               "product": "Vegan Pea Protein", **SPEND}
    sender = {**SENDER, "name": "Aditi", "role": "Partnerships"}
    d = drafting.TemplateDrafter().initial(contact, sender, DEFAULTS["pitches"], MARKET)
    assert d.subject == "GreenLeaf Nutrition, your marketing bill for ₹25 lakhs is ready"
    assert d.body == (
        "Hi GreenLeaf Nutrition team,\n\nIn the last quarter, your estimated range of spend on marketing was ₹18–25 lakhs.\n\n"
        "We don't know how many customers bought Vegan Pea Protein after watching someone dance with it online.\n\n"
        "[SCREENSHOT_1]\n\nBut we can take your product directly into the inbox of your target consumer — the shopper whose diet, taste and "
        "budget already match it.\n\n[SCREENSHOT_2]\n\n[SCREENSHOT_3]\n\nSet up a call with Dr. Tracker Bharat: https://cal.example/drtracker\n\n"
        "Aditi\nPartnerships, Dr. Tracker Bharat")
    on = drafting.TemplateDrafter().initial(contact, sender, DEFAULTS["pitches"], {**MARKET, "mention_trial": True})
    assert MARKET["trial_text"] not in on.body  # the guide has no trial
    assert not qa.has_errors(qa.check(d.subject, d.body, contact, sender, DEFAULTS["pitches"], "initial", MARKET))
    no_hero = drafting.TemplateDrafter().initial({**contact, "product": "Plant-Based Nutrition"}, sender, DEFAULTS["pitches"], MARKET)
    assert "bought your product after watching" in no_hero.body
    assert not qa.has_errors(qa.check(no_hero.subject, no_hero.body, contact, sender, DEFAULTS["pitches"], "initial", MARKET))
    # A figure that isn't the sheet's is still blocked.
    assert qa.has_errors(qa.check(d.subject, d.body.replace("₹18–25", "₹30–40"), contact, sender, DEFAULTS["pitches"], "initial", MARKET))


@pytest.mark.parametrize("change,reason", [({"spend": ""}, "no spend figure"), ({"spend": "a lot"}, "no spend figure")])
def test_keyless_brand_template_stops_without_spend(change, reason):
    with pytest.raises(drafting.DraftError, match=reason) as e:
        drafting.TemplateDrafter().initial({**BRAND, **SPEND, **change}, SENDER, DEFAULTS["pitches"], MARKET)
    assert e.value.skip


def test_keyless_brand_template_stops_without_a_call_link():
    with pytest.raises(drafting.DraftError, match="no call link") as e:
        drafting.TemplateDrafter().initial({**BRAND, **SPEND}, SENDER, DEFAULTS["pitches"], {**MARKET, "meeting_link": ""})
    assert e.value.skip


def test_keyless_brand_followups_are_the_guides_touches_2_and_3():
    contact, prev = {**BRAND, **SPEND}, {"subject": "Heka Bites, your marketing bill for ₹25 lakhs is ready", "message_id": None}
    t2 = drafting.TemplateDrafter().followup(contact, SENDER, DEFAULTS["pitches"], prev, 1, MARKET)
    assert t2.subject == prev["subject"]
    assert t2.body == ("Hi Heka Bites team,\n\nQuick nudge on this. One question to think about: of last quarter's ₹18–25 lakhs, "
                       "how much can you tie to an actual order?\n\nWith us, every sale comes through a code. 15 minutes to "
                       "show you: https://cal.example/drtracker\n\nAnanya Gupta")
    t3 = drafting.TemplateDrafter().followup(contact, SENDER, DEFAULTS["pitches"], prev, 2, MARKET)
    assert t3.body.startswith("Hi Heka Bites team,\n\nLast note from me.") and "I won't follow up again." in t3.body
    for d in (t2, t3):
        assert not qa.has_errors(qa.check(d.subject, d.body, contact, SENDER, DEFAULTS["pitches"], "followup", MARKET))
    with pytest.raises(drafting.DraftError, match="never sends more than 3") as e:
        drafting.TemplateDrafter().followup(contact, SENDER, DEFAULTS["pitches"], prev, 3, MARKET)
    assert e.value.skip


def _fake_claude(payload):
    seen = {}

    def create(**kwargs):
        seen.update(kwargs)
        return SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text=json.dumps(payload))])

    d = drafting.ClaudeDrafter("sk-test")
    d.client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(create=create)))
    return d, seen


def test_claude_writes_brands_to_the_guide_without_pricing_or_phone():
    d, seen = _fake_claude({"skip_reason": "", "subject": "Straight into your buyer's chat", "body": "Hi Heka Bites team,\n\nBody."})
    contact = {**BRAND, "product": "Makhana", "pricing": "₹199–₹499", "phone": "98765 43210",
               "extra": json.dumps({"Founded": "2019", "Alt phone": "12345"}), "site_text": "Title: Heka"}
    out = d.initial({**contact, **SPEND}, SENDER, DEFAULTS["pitches"], {**MARKET, "meeting_link": "https://cal.example/x"})
    assert out.subject == "Straight into your buyer's chat"
    assert "Brand Outreach Agent" in seen["system"] and "[SCREENSHOT_1]" in seen["system"] and "[SCREENSHOT_3]" in seen["system"]
    user = seen["messages"][0]["content"]
    assert "touch 1" in user and '"hero_product": "Makhana"' in user and "Founded: 2019" in user
    assert '"call_link": "https://cal.example/x"' in user and '"spend_low_lakhs": 18' in user and '"spend_high_lakhs": 25' in user
    assert "₹199" not in user and "98765" not in user and "12345" not in user
    assert '"sender_name": "Ananya Gupta"' in user and '"sender_title": "Founder"' in user and "<website>\nTitle: Heka" in user
    assert seen["output_config"]["format"]["schema"] == drafting.GUIDED_SCHEMA


def test_claude_writes_brand_followups_from_the_guides_sequence():
    d, seen = _fake_claude({"skip_reason": "", "subject": "ignored", "body": "Hi Heka Bites team,\n\nQuick nudge on this."})
    prev = {"subject": "Heka Bites, your marketing bill for ₹25 lakhs is ready", "body": "Hi Heka Bites team,\n\n…"}
    out = d.followup({**BRAND, **SPEND}, SENDER, DEFAULTS["pitches"], prev, 1, MARKET)
    assert out.subject == prev["subject"] and "Brand Outreach Agent" in seen["system"]
    assert "touch 2" in seen["messages"][0]["content"] and '"touch_number": 2' in seen["messages"][0]["content"]


def test_a_row_the_guide_skips_is_flagged_not_retried_and_can_be_forced(configured, monkeypatch):
    c = configured
    cid = c.post("/api/contacts", json={"name": "Medicine Nest", "email": "hi@nest.test", "type": "brand", "category": "SUP"}).json()["id"]
    other = c.post("/api/contacts", json={"name": "OZiva", "email": "care@oziva.test", "type": "brand", "category": "SUP"}).json()["id"]
    monkeypatch.setattr(research, "research", lambda contact: ("", "", []))

    class Skipper:
        name = "claude"

        def initial(self, contact, sender, pitches, market=None):
            if contact["name"] == "Medicine Nest" and not contact.get("_write_anyway"):
                raise drafting.DraftError("Not written: a store or reseller, not a brand.", skip=True)
            return drafting.Draft("Subject line", "Hi " + contact["name"] + " team,\n\nBody.", "claude")

    assert worker.draft_stage(drafter=Skipper()) == 1  # the skip doesn't stop the next contact being drafted
    with db.connect() as conn:
        nest, oz = outreach.get_contact(conn, cid), outreach.get_contact(conn, other)
    assert nest["last_error"].startswith("Not written:") and nest["draft_attempts"] == outreach.MAX_DRAFT_ATTEMPTS
    assert oz["status"] == "drafted"
    assert worker.draft_stage(drafter=Skipper()) == 0  # not tried again
    m = outreach.draft(cid, drafter=Skipper(), force=True)
    assert m["status"] == "draft"


# --- Research: website pages and socials --------------------------------------------------

HOME = """<html><head><title>Heka Bites</title><meta name="description" content="Roasted makhana snacks."></head>
<body><nav><a href="/pages/about-us">About</a><a href="/collections/all">Shop</a><a href="https://other.test/about">x</a></nav>
<h1>Snacks that crunch</h1><p>Made in Hyderabad.</p></body></html>"""
ABOUT = "<html><head><title>About Heka</title></head><body><h1>Our story</h1><p>Started in a home kitchen in 2019.</p></body></html>"
SHOP = "<html><head><title>Shop Heka</title></head><body><h1>All products</h1><p>Peri peri makhana, cheese makhana.</p></body></html>"
INSTA = '<html><head><meta property="og:title" content="Heka Bites (@hekabites)"><meta property="og:description" content="Roasted makhana from Hyderabad."></head></html>'
LOGIN = '<html><head><title>Log into Facebook</title><meta name="description" content="Log into Facebook to start sharing."></head></html>'


def test_social_urls_from_the_sheet():
    c = {"website": "heka.test", "socials": "@hekabites; https://in.linkedin.com/company/heka and facebook.com/heka",
         "extra": json.dumps({"Press": "https://news.test/heka-story", "Founded": "2019"})}
    assert research.social_urls(c) == ["https://in.linkedin.com/company/heka", "https://facebook.com/heka",
                                       "https://www.instagram.com/hekabites/"]
    assert research.social_urls({"socials": "Not publicly listed"}) == []
    assert research.social_urls({"socials": "", "extra": json.dumps({"Press": "https://news.test/heka"})}) == ["https://news.test/heka"]


def test_research_reads_inner_pages_and_socials_and_notes_what_it_couldnt(monkeypatch):
    pages = {"https://heka.test/": HOME, "https://heka.test/pages/about-us": ABOUT, "https://heka.test/collections/all": SHOP,
             "https://www.instagram.com/hekabites/": INSTA, "https://facebook.com/heka": LOGIN}

    def fake(url):
        if url not in pages:
            raise research.ResearchError("The site answered 999")
        return pages[url]

    monkeypatch.setattr(research, "fetch_html", fake)
    text, read, missed = research.research({"website": "heka.test", "socials": "@hekabites facebook.com/heka linkedin.com/company/heka"})
    assert "[Website https://heka.test/]" in text and "Started in a home kitchen" in text and "cheese makhana" in text
    assert "[Instagram https://www.instagram.com/hekabites/]" in text and "Roasted makhana from Hyderabad." in text
    assert "other.test" not in read and "Log into" not in text
    assert any(m.startswith("Facebook: shows a login page") for m in missed) and any(m.startswith("LinkedIn:") for m in missed)


def test_prompt_marks_socials_as_their_marketing():
    contact = {**BRAND, "site_text": "[Instagram x]\nDescription: 12K followers"}
    text = prompts.guided_initial(contact, SENDER, MARKET)
    assert "follower counts" in text and "<website>" in text


# --- The brand guide's stop rules, cleaner names and products, and estimated spend ------------------------

@pytest.mark.parametrize("email,skipped", [
    ("support@heka.test", True), ("customercare@heka.test", True), ("care@heka.test", True), ("complaints@heka.test", True),
    ("feedback@heka.test", True), ("noreply@heka.test", True),
    ("info@heka.test", False), ("hello@heka.test", False), ("contact@heka.test", False), ("careers@heka.test", False),
    ("founder@heka.test", False), ("heka.foods@gmail.com", False)])
def test_support_inboxes_are_skipped(email, skipped):
    contact = {**BRAND, **SPEND, "email": email}
    if skipped:
        with pytest.raises(drafting.DraftError, match="support or complaints inbox") as e:
            drafting.TemplateDrafter().initial(contact, SENDER, DEFAULTS["pitches"], MARKET)
        assert e.value.skip
        # A person can still ask for it.
        drafting.TemplateDrafter().initial({**contact, "_write_anyway": True}, SENDER, DEFAULTS["pitches"], MARKET)
    else:
        drafting.TemplateDrafter().initial(contact, SENDER, DEFAULTS["pitches"], MARKET)


@pytest.mark.parametrize("change", [{"name": "Bira 91 Craft Beer"}, {"product": "Nicotine-free hookah flavours"},
                                    {"product": "Weight loss shakes for kids"}, {"product": "Diabetes Cure Tea"}])
def test_brands_the_guide_says_are_not_a_fit_are_skipped(change):
    with pytest.raises(drafting.DraftError, match="section 7") as e:
        drafting.TemplateDrafter().initial({**BRAND, **SPEND, **change}, SENDER, DEFAULTS["pitches"], MARKET)
    assert e.value.skip


def test_claude_is_not_called_for_a_row_the_stop_rules_skip():
    d, seen = _fake_claude({"skip_reason": "", "subject": "s", "body": "b"})
    with pytest.raises(drafting.DraftError, match="support or complaints inbox"):
        d.initial({**BRAND, **SPEND, "email": "support@heka.test"}, SENDER, DEFAULTS["pitches"], MARKET)
    assert not seen


@pytest.mark.parametrize("cell,hero", [
    ("ABC Millet Mix / Multi Millet Chikki", "ABC Millet Mix"), ("SanketLife ECG; EasyTouch", "SanketLife ECG"),
    ("Raw Unpasteurized A2 Gir Cow Milk", "Raw Unpasteurized A2 Gir Cow Milk"),
    ("Nutraceutical contract development and manufacturing", ""), ("Ayurvedic botanical skin and hair care", ""),
    ("Organic staples, superfoods, seeds, spices", ""), ("Build Mass Gainer + 4 free items combo", ""),
    ("Healthcare and nutraceutical solutions", ""), ("Sports nutrition", "")])
def test_hero_product_is_a_product_not_a_description_or_an_offer(cell, hero):
    assert drafting.hero_product(cell, "") == hero


@pytest.mark.parametrize("name,shown", [
    ("Natures Basket - Sultan Palya", "Natures Basket"), ("Fix The Glytch / GLYTCH", "Fix The Glytch"),
    ("Fast&Up Mumbai", "Fast&Up"), ("Primal Nutrition.In", "Primal Nutrition"), ("Desi Cow Milk Hyd", "Desi Cow Milk"),
    ("Mediwheel - Your Wellness Partner", "Mediwheel"), ("Agatsa (Bengaluru Hub)", "Agatsa"),
    ("Alphacentric Healthcare Pvt Ltd", "Alphacentric Healthcare")])
def test_brand_names_are_the_names_shoppers_know(name, shown):
    assert drafting.brand_name(name) == shown


def test_same_brand_in_another_city_tab_is_a_duplicate():
    assert importer.brand_key("Fast&Up Mumbai") == importer.brand_key("Fast&Up (Aeronutrix Sports Products)")
    assert importer.brand_key("Natures Basket - Bellandur") == importer.brand_key("Natures Basket - Garuda Mall")


def test_spend_estimate_follows_the_brands_size():
    tiny = spend.estimate({"email": "millets@gmail.com", "category": "FD", "product": "Millet Mix"})
    typical_row = {"email": "hi@brand.test", "website": "brand.test", "socials": "https://instagram.com/brand",
                   "category": "PRO", "product": "Whey Protein"}
    typical = spend.estimate(typical_row)
    big = spend.estimate({**{"email": "hi@brand.test", "website": "brand.test", "socials": "@brand", "category": "PRO",
                             "product": "Bars; Shakes; Muesli", "pricing": "₹530–₹1,999"},
                          "site_text": "Order on Blinkit, Zepto and Amazon. As seen on Shark Tank. In 2,000+ stores."})
    assert tiny[:2] == (0.5, 1.5) and "free-mail" in tiny[2]
    assert typical[:2] == (3, 8)
    assert big[:2] == (20, 50) and "Shark Tank" in big[2] and big[2].startswith(spend.SOURCE_PREFIX)
    insta = "[Instagram https://www.instagram.com/brand/]\nDescription: 250K Followers, 10 Following, 900 Posts"
    assert spend.estimate({"email": "a@b.test", "site_text": insta})[:2] == (10, 25)
    # LinkedIn's follower count (often a few hundred) is not Instagram's, and never pulls an estimate down.
    linkedin = "[LinkedIn https://in.linkedin.com/company/brand]\nDescription: Brand | 139 followers on LinkedIn."
    assert spend.followers(linkedin) is None
    assert spend.estimate({**typical_row, "site_text": linkedin})[:2] == typical[:2]


def test_estimated_spend_is_saved_and_every_touch_quotes_it(configured, monkeypatch):
    c = configured
    monkeypatch.setattr(research, "research", lambda contact: ("Buy on Blinkit and Amazon.", "https://heka.test/", []))
    cid = c.post("/api/contacts", json={"name": "Heka Bites", "email": "hello@heka.test", "type": "brand", "category": "FD",
                                         "website": "heka.test", "socials": "@hekabites", "product": "Makhana"}).json()["id"]
    m = c.post(f"/api/contacts/{cid}/draft").json()
    contact = contacts(c)["hello@heka.test"]
    low, high = marketplace.spend_range(contact["spend"])
    assert contact["spend_source"].startswith(spend.SOURCE_PREFIX) and "quick-commerce" in contact["spend_source"]
    figure = f"₹{marketplace.lakhs(low)}–{marketplace.lakhs(high)} lakhs"
    assert f"your estimated range of spend on marketing was {figure}" in m["body"] and not qa.has_errors(m["qa"]), m["qa"]
    c.post(f"/api/messages/{m['id']}/approve")
    outreach.send(m["id"], transport=Capture())
    with db.connect() as conn:
        conn.execute("UPDATE contacts SET last_sent_at = 0 WHERE id = ?", (cid,))
    follow = c.post(f"/api/contacts/{cid}/draft?kind=followup").json()
    assert figure in follow["body"] and contacts(c)["hello@heka.test"]["spend"] == contact["spend"]


def test_a_spend_figure_from_the_sheet_is_never_replaced(configured):
    c = configured
    cid = c.post("/api/contacts", json={"name": "Heka Bites", "email": "hello@heka.test", "type": "brand", "category": "FD",
                                         "spend": "18-25", "spend_source": "Founder"}).json()["id"]
    c.post(f"/api/contacts/{cid}/draft")
    assert (contacts(c)["hello@heka.test"]["spend"], contacts(c)["hello@heka.test"]["spend_source"]) == ("18-25", "Founder")


@pytest.mark.parametrize("name", ["Jaslok Hospital Outpatient Clinic", "Qi Spine Clinic - Churchgate", "V-Rehab Physiotherapy Clinic",
                                  "Waves Gym Personal Training", "Shathayu Ayurveda Wellness Centre", "Dr L H Hiranandani Hospital OP"])
def test_clinics_and_hospitals_in_the_brand_sheet_are_skipped(name):
    with pytest.raises(drafting.DraftError, match="a service") as e:
        drafting.TemplateDrafter().initial({**BRAND, **SPEND, "name": name}, SENDER, DEFAULTS["pitches"], MARKET)
    assert e.value.skip


@pytest.mark.parametrize("name", ["Tulip LAB Pvt. LTD", "Gym Bazar", "Optimum Nutrition", "Clinikally"])
def test_product_makers_are_not_mistaken_for_services(name):
    drafting.TemplateDrafter().initial({**BRAND, **SPEND, "name": name}, SENDER, DEFAULTS["pitches"], MARKET)


@pytest.mark.parametrize("name", ["Miracle Drinks In", "Fit Treat Nutrition"])
def test_a_brands_own_name_is_not_a_claim(name):
    contact = {**BRAND, **SPEND, "name": name}
    d = drafting.TemplateDrafter().initial(contact, SENDER, DEFAULTS["pitches"], MARKET)
    issues = qa.check(d.subject, d.body, contact, SENDER, DEFAULTS["pitches"], "initial", MARKET)
    assert not issues, issues
    assert qa.has_errors(qa.check(d.subject, d.body + "\n\nIt's a miracle.", contact, SENDER, DEFAULTS["pitches"], "initial", MARKET))


def test_well_known_brands_use_the_reference_file():
    low, high, basis = spend.estimate({"name": "boAt (Bengaluru Branch)", "email": "a@boat.test"})
    assert (low, high) == (2000, 5000) and "consumer-electronics" in basis
    contact = {"name": "boAt (Bengaluru Branch)", "email": "partners@boat.test", "type": "brand", "category": "DV",
               "spend": spend.as_cell(low, high), "spend_source": basis}
    d = drafting.TemplateDrafter().initial(contact, SENDER, DEFAULTS["pitches"], MARKET)
    assert d.subject == "boAt, your marketing bill for ₹5,000 lakhs is ready"
    assert "estimated range of spend on marketing was ₹2,000–5,000 lakhs" in d.body
    assert not qa.has_errors(qa.check(d.subject, d.body, contact, SENDER, DEFAULTS["pitches"], "initial", MARKET))


@pytest.mark.parametrize("name,reason", [
    ("Rujuta Diwekar", "a professional"), ("MedPlus Hyderabad", "a store or reseller"), ("Whey Protein Store", "a store or reseller"),
    ("iCALL Psychosocial Helpline (TISS)", "a service"), ("Inner Space Counseling", "a service"),
    ("Clear Skin Centre - Bandra", "a service")])
def test_people_stores_and_services_in_the_brand_sheet_are_skipped(name, reason):
    with pytest.raises(drafting.DraftError, match=reason) as e:
        drafting.TemplateDrafter().initial({**BRAND, **SPEND, "name": name}, SENDER, DEFAULTS["pitches"], MARKET)
    assert e.value.skip


# --- Screenshots inside the body: one marker per screenshot ---------------------------------------

@pytest.mark.parametrize("body,layout", [
    ("A\n\n[SCREENSHOT_1]\n\nB\n\n[SCREENSHOT_2]\n\n[SCREENSHOT_3]\n\nC", ["A", 0, "B", 1, 2, "C"]),
    ("A\n\n[HOW_IT_WORKS_IMAGE]\n\nB", ["A", 0, 1, 2, "B"]),                       # one marker: all of them there
    ("A\n\n[SCREENSHOT_1]\n\nB\n\n[SCREENSHOT_2]\n\nC", ["A", 0, "B", 1, 2, "C"]),  # a missing one follows the last
    ("A\n\n[screenshot_2]\n\nB\n\n[SCREENSHOT_1]\n\nC", ["A", 1, "B", 0, 2, "C"]),
])
def test_screenshots_go_where_their_markers_are(body, layout):
    got = templates.segments(body, 3)
    assert [p.strip() if isinstance(p, str) else p for p in got] == layout


def test_qa_checks_the_screenshot_markers():
    contact = {**BRAND, **SPEND}
    d = drafting.TemplateDrafter().initial(contact, SENDER, DEFAULTS["pitches"], MARKET)
    run = lambda body: qa.check(d.subject, body, contact, SENDER, DEFAULTS["pitches"], "initial", MARKET)
    assert not run(d.body)
    assert qa.has_errors(run(d.body.replace("[SCREENSHOT_3]", "[SCREENSHOT_2]")))  # twice
    assert qa.has_errors(run(d.body.replace("[SCREENSHOT_3]", "[SCREENSHOT_4]")))  # no such image
    missing = run(d.body.replace("[SCREENSHOT_3]\n\n", ""))
    assert not qa.has_errors(missing) and any("[SCREENSHOT_3]" in i["message"] for i in missing)


def test_preview_lays_out_text_and_screenshots_like_the_email(configured):
    c = configured
    cid = c.post("/api/contacts", json={"name": "Heka Bites", "email": "hello@heka.test", "type": "brand", "category": "FD",
                                         **SPEND}).json()["id"]
    m = c.post(f"/api/contacts/{cid}/draft").json()
    parts = c.get(f"/api/preview/{m['id']}").json()["parts"]
    kinds = ["image" if "image" in p else "text" for p in parts]
    assert kinds == ["text", "image", "text", "image", "image", "text"]
    assert [p["image"] for p in parts if "image" in p] == [f"/api/templates/brand/image/{i}" for i in range(3)]
    assert "unsubscribe" in parts[-1]["text"] and "Set up a call" in parts[-1]["text"]
