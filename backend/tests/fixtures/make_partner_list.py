"""Builds partner_list_sample.xlsx: a synthetic copy of the founder's workbook layout.

Every name, email and site is invented. Run it again after changing a row:
    .venv/bin/python backend/tests/fixtures/make_partner_list.py
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).with_name("partner_list_sample.xlsx")
HEADERS = ["ID", "Name", "City", "Type", "Category", "Tier", "Phone", "Email", "Status", "Notes / Source"]

# (tab, row values). Comments say which edge case each row covers.
ROWS = {
    "Hyderabad": [
        ["B-FD-T2-HYD-01", "Crunchwell Foods", "HYD", "B", "FD", "T2", "+91 90000 00001", "hello@crunchwell.test", "",
         "Official site crunchwell.test; Millet snack bars and roasted makhana with no added sugar, Hyderabad HQ; Source: https://crunchwell.test/about"],
        # person in the name, explicit "Not contacted"
        ["P-NT-T1-HYD-01", "Dr. Meera Kulkarni (FitRoot)", "HYD", "P", "NT", "T1", "", "meera@fitroot.test", "Not contacted",
         "Official site fitroot.test; Clinical nutritionist running weight-management programmes; Source: https://fitroot.test"],
        # government helpline by name -> never imported
        ["P-MH-TBD-HYD-01", "Tele-MANAS Hyderabad", "HYD", "P", "MH", "TBD", "14416", "telemanas@helpline.test", "",
         "Government mental health line; Source: https://telemanas.test"],
        # already contacted by hand, generic inbox, CHECK note
        ["B-SUP-T3-HYD-01", "Sunleaf Naturals", "HYD", "B", "SUP", "T3", "", "info@sunleaf.test", "Contacted",
         "Plant-based vitamin gummies; CHECK: generic inbox, find a named contact"],
    ],
    "Delhi": [
        ["P-GP-T2-DEL-01", "Dr. Arvind Bhalla (Bhalla Family Clinic)", "DEL", "P", "GP", "T2", "", "arvind@bhallaclinic.test", "",
         "Family physician, general OPD and preventive check-ups, Delhi"],
        # organisation only, replied
        ["P-SPC-T1-DEL-01", "Shreya Eye Clinic", "DEL", "P", "SPC", "T1", "", "care@shreyaclinic.test", "Replied",
         "Official site shreyaclinic.test; Cataract and retina clinic; Source: https://shreyaclinic.test"],
        # no email
        ["B-DV-T2-DEL-01", "VitalSense", "DEL", "B", "DV", "T2", "+91 90000 00007", "", "", "Home blood pressure monitors"],
        # helpline matched by name
        ["P-MH-T2-DEL-01", "iCALL Psychosocial Support", "DEL", "P", "MH", "T2", "", "icall@support.test", "", ""],
        # invalid email
        ["B-PRO-T1-DEL-01", "LiftFuel Nutrition", "DEL", "B", "PRO", "T1", "", "sales@liftfuel", "", "Whey protein and bars"],
    ],
    "Mumbai": [
        # TBD tier (the ID changes when it's filled in), mixed-case generic email, marketing claim in the notes
        ["B-SUP-TBD-MUM-01", "Greenbowl Supplements", "MUM", "B", "SUP", "TBD", "", "Support@Greenbowl.test", "Interested",
         "Official site greenbowl.test; India's No. 1 ashwagandha capsules"],
        # the "Dr. Sports" name, unsubscribed
        ["P-PT-T2-MUM-01", "Total Health Physio & Sports", "MUM", "P", "PT", "T2", "", "info@totalhealthphysio.test", "Unsubscribed",
         "Sports physiotherapy and rehab, Mumbai HQ"],
        # Dt. title, empty tier, empty notes
        ["P-NT-T3-MUM-01", "Dt. Kavya Rao (Greenplate)", "MUM", "P", "NT", "", "", "kavya@greenplate.test", "", ""],
        # duplicate of the Hyderabad Crunchwell email (different case) -> skipped
        ["B-FD-T2-MUM-02", "Crunchwell Foods Mumbai", "MUM", "B", "FD", "T2", "", "HELLO@crunchwell.test", "", "Mumbai stockist"],
    ],
    "Bengaluru": [
        # bounced -> imported and suppressed
        ["B-PRO-T2-BLR-01", "FitPlate Protein", "BLR", "B", "PRO", "T2", "", "team@fitplate.test", "Bounced",
         "Whey and plant protein tubs; Source: https://fitplate.test"],
        # person without a title
        ["P-MH-T1-BLR-01", "Ritika Sen (CalmCorner)", "BLR", "P", "MH", "T1", "", "ritika@calmcorner.test", "",
         "Counselling psychologist offering online therapy for anxiety and stress"],
        # helpline matched only by the notes
        ["P-MH-TBD-BLR-02", "Sahara Support Line", "BLR", "P", "MH", "TBD", "", "sahara@support.test", "",
         "24x7 suicide prevention line run by volunteers"],
        # a Status the app doesn't recognise -> skipped so nobody is contacted twice by mistake
        ["P-NT-T2-BLR-02", "NutriNest", "BLR", "P", "NT", "T2", "", "hello@nutrinest.test", "Called twice", "Diet plans for PCOS"],
    ],
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
STATUS_FILL = PatternFill("solid", fgColor="FFF2CC")


def build(path: Path = OUT) -> Path:
    wb = openpyxl.Workbook()
    legend = wb.active
    legend.title = "Legend"
    for row in [["Dr.Tracker partner list: legend (synthetic test data)"], [], ["Code", "Meaning"],
                ["B", "Brand"], ["P", "Professional"], ["FD / SUP / PRO / DV", "Food / Supplements / Protein / Devices"],
                ["NT / GP / SPC / MH / PT", "Nutritionist / GP / Specialist / Mental health / Physio-trainer"],
                ["T1 / T2 / T3 / TBD", "Priority tier"]]:
        legend.append(row)
    legend["A1"].font = Font(bold=True, size=14)

    for tab, rows in ROWS.items():
        ws = wb.create_sheet(tab)
        ws.append([f"{tab}: partner outreach list (synthetic test data — not real partners)"])
        ws.append(["Every name, email and site here is invented. Status is updated by the outreach agent."])
        ws.append([])
        ws.append(HEADERS)
        ws.merge_cells("A1:J1")
        ws["A1"].font = Font(bold=True, size=13)
        for cell in ws[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
        for r in rows:
            ws.append(r)
            ws.cell(ws.max_row, 9).fill = STATUS_FILL
        # trailing empty but formatted rows, like a real sheet
        for i in range(3):
            ws.cell(ws.max_row + 1, 9).fill = STATUS_FILL
        ws.freeze_panes = "A5"
        for col, width in zip("ABCDEFGHIJ", (18, 38, 6, 6, 9, 6, 16, 30, 14, 70)):
            ws.column_dimensions[col].width = width
    wb.save(path)
    return path


if __name__ == "__main__":
    print(build())
