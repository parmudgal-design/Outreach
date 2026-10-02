"""Status back to the founder's workbook: only Status cells change."""
import csv
import io
import json

import openpyxl

from backend.app import db
from backend.app.services import outreach
from backend.tests.test_partner_sheet import FIXTURE, XLSX, commit, contacts, edited, upload

STATUS_COL = 9  # I


def export(c, data: bytes):
    r = c.post("/api/export/partner-workbook", files={"file": ("partners.xlsx", data, XLSX)})
    assert r.status_code == 200, r.text
    return openpyxl.load_workbook(io.BytesIO(r.content)), json.loads(r.headers["x-status-summary"])


def set_status(email: str, status: str):
    with db.connect() as conn:
        cid = conn.execute("SELECT id FROM contacts WHERE email = ?", (email,)).fetchone()[0]
        outreach.set_status(conn, cid, status)


def test_status_written_back_and_nothing_else_changes(configured):
    c = configured
    commit(c, upload(c))
    set_status("hello@crunchwell.test", "sent")
    set_status("meera@fitroot.test", "replied")
    set_status("arvind@bhallaclinic.test", "adopted")
    set_status("kavya@greenplate.test", "declined")
    set_status("ritika@calmcorner.test", "bounced")

    original = openpyxl.load_workbook(FIXTURE)
    wb, summary = export(c, FIXTURE.read_bytes())
    assert summary["updated"] == 6  # the 5 above, plus Mumbai's duplicate row for the same Crunchwell address
    assert summary["unmatched"] == 6  # 3 helplines, no email, invalid email, unknown status: never imported

    assert wb.sheetnames == original.sheetnames  # Legend kept, same order
    for ws in original.worksheets:
        out = wb[ws.title]
        assert out.max_row == ws.max_row and out.max_column == ws.max_column
        assert out.merged_cells.ranges == ws.merged_cells.ranges and out.freeze_panes == ws.freeze_panes
        for row in ws.iter_rows():
            for cell in row:
                new = out.cell(cell.row, cell.column)
                if ws.title != "Legend" and cell.column == STATUS_COL and cell.row > 4:
                    continue
                assert new.value == cell.value, (ws.title, cell.coordinate)
                assert new.fill.fgColor.rgb == cell.fill.fgColor.rgb and new.font.b == cell.font.b, (ws.title, cell.coordinate)

    status = lambda tab, row: wb[tab].cell(row, STATUS_COL).value
    assert status("Hyderabad", 5) == "Contacted"
    assert status("Hyderabad", 6) == "Replied"
    assert status("Delhi", 5) == "Interested"
    assert status("Mumbai", 7) == "Not interested"
    assert status("Bengaluru", 6) == "Bounced"
    assert status("Mumbai", 8) == "Contacted"              # skipped as a duplicate, but it's the same address
    assert status("Hyderabad", 8) == "Contacted"           # already said so: untouched
    assert status("Mumbai", 6) == "Unsubscribed"
    assert status("Hyderabad", 7) in ("", None)            # helpline row: never imported, left alone
    assert status("Bengaluru", 8) == "Called twice"        # unknown status row: left alone
    assert wb["Hyderabad"].cell(9, STATUS_COL).fill.fgColor.rgb == original["Hyderabad"].cell(9, STATUS_COL).fill.fgColor.rgb


def test_matches_by_partner_id_when_the_email_changed(configured):
    c = configured
    commit(c, upload(c))
    set_status("support@greenbowl.test", "declined")
    # The founder later fixed the tier and the email in the sheet.
    data = edited({("Mumbai", 5, "A"): "B-SUP-T2-MUM-01", ("Mumbai", 5, "H"): "partners@greenbowl.test"})
    wb, _ = export(c, data)
    assert wb["Mumbai"].cell(5, STATUS_COL).value == "Not interested"
    assert wb["Mumbai"].cell(5, 8).value == "partners@greenbowl.test"  # the sheet's email isn't touched


def test_drafted_counts_as_not_contacted(configured):
    c = configured
    commit(c, upload(c))
    cid = contacts(c)["hello@crunchwell.test"]["id"]
    c.post(f"/api/contacts/{cid}/draft")
    _, summary = export(c, FIXTURE.read_bytes())
    assert summary["updated"] == 0  # blank Status cells stay blank


def test_rejects_a_workbook_without_partner_tabs(app_client):
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    r = app_client.post("/api/export/partner-workbook", files={"file": ("other.xlsx", buf.getvalue(), XLSX)})
    assert r.status_code == 400 and "partner tabs" in r.json()["detail"]


def test_status_csv_keyed_by_partner_id(configured):
    c = configured
    commit(c, upload(c))
    set_status("meera@fitroot.test", "adopted")
    rows = list(csv.DictReader(io.StringIO(c.get("/api/export/partner-status.csv").text)))
    by_id = {r["partner_id"]: r for r in rows}
    assert len(rows) == 10
    assert by_id["P-NT-T1-HYD-01"]["status"] == "Interested" and by_id["P-NT-T1-HYD-01"]["app_status"] == "adopted"
    assert by_id["B-SUP-T3-HYD-01"]["status"] == "Contacted"
    assert by_id["P-PT-T2-MUM-01"]["status"] == "Unsubscribed"
